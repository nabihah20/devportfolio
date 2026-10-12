import docx, os, sys
import win32com.client
import fitz
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml

PRIMARY_COLOR = '2F5496' # Elegant Professional Navy Blue
TEXT_COLOR = '111111'
SUBTEXT_COLOR = '444444'

def create_resume_2page(output_docx):
    doc = docx.Document()
    
    # Page setup: A4, 0.65" margins
    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.top_margin = Inches(0.65)
        s.bottom_margin = Inches(0.65)
        s.left_margin = Inches(0.65)
        s.right_margin = Inches(0.65)
        
        # Add footer with page numbering
        footer = s.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ft.paragraph_format.space_before = Pt(4)
        r_ft = p_ft.add_run("Nabihah Binti Taquddin Azmy — Resume")
        r_ft.font.name = 'Arial'
        r_ft.font.size = Pt(8.0)
        r_ft.font.color.rgb = RGBColor.from_string('777777')

    base_pt = 9.5
    line_sp = 1.15

    def add_p(text='', bold=False, italic=False, size=base_pt, color=TEXT_COLOR, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=3, line_spacing=line_sp):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.italic = italic
            r.font.name = 'Arial'
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        return p

    def add_sec_hdr(title, space_before=7, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(title)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor.from_string(PRIMARY_COLOR)
        
        pPr = p._p.get_or_add_pPr()
        border_xml = f'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="6" w:space="2" w:color="{PRIMARY_COLOR}"/></w:pBdr>'
        pPr.append(parse_xml(border_xml))
        return p

    def add_bullet(prefix, text, space_after=2.5, level=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_sp
        indent = 0.2 if level == 0 else 0.35
        p.paragraph_format.left_indent = Inches(indent)
        
        bullet_char = '• ' if level == 0 else '○ '
        rb = p.add_run(bullet_char)
        rb.font.name = 'Arial'
        rb.font.size = Pt(base_pt)
        rb.font.color.rgb = RGBColor.from_string(PRIMARY_COLOR)
        
        if prefix:
            rp = p.add_run(prefix)
            rp.bold = True
            rp.font.name = 'Arial'
            rp.font.size = Pt(base_pt)
            rp.font.color.rgb = RGBColor.from_string('000000')
            
        rt = p.add_run(text)
        rt.bold = False
        rt.font.name = 'Arial'
        rt.font.size = Pt(base_pt)
        rt.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
        return p

    # ================= PAGE 1 =================
    # Header
    add_p('NABIHAH BINTI TAQUDDIN AZMY', bold=True, size=15.0, color='000000', align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5)
    add_p('Software Engineer | Embedded & Full-Stack Systems', bold=True, size=10.5, color=PRIMARY_COLOR, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2.5)
    add_p('Perai, Penang | +60 13-519 8421 | nabihahtaqudin@gmail.com | linkedin.com/in/nabihah-taquddin-azmy', bold=False, size=9.0, color='444444', align=WD_ALIGN_PARAGRAPH.CENTER, space_after=5)

    # Professional Summary
    add_sec_hdr('PROFESSIONAL SUMMARY', space_before=4)
    add_p('Skilled Software Engineer and Technical Lead with over 5 years of professional experience building full-stack web applications, relational database infrastructure, and leading software platform delivery. Combines a strong software engineering foundation (Node.js, Express, NestJS, RESTful APIs, MySQL, MongoDB) with hands-on capabilities in embedded hardware integration (ESP32-S3 microcontrollers, Embedded Touchscreen GUIs, BLE telemetry, Edge AI). Experienced lead Person-in-Charge (PIC) for multi-module enterprise platforms, managing end-to-end SDLC workflows, Functional Specification Documentation (FSD), system monitoring, and cross-functional project execution.', size=base_pt, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4)

    # Core Competencies & Tech Stack
    add_sec_hdr('CORE COMPETENCIES & TECH STACK', space_before=5)
    add_bullet('System Architecture & Integration: ', 'End-to-End System Design, Hardware-Software Integration, Embedded Touchscreen GUI Architecture, RESTful API Design, System Process Mapping, FSD Authoring & Technical Documentation.')
    add_bullet('IDEs, Embedded UI & Hardware Frameworks: ', 'ESP-IDF (C/C++), Arduino IDE, SquareLine Studio (LVGL Framework), Microcontrollers (ESP32-S3, Arduino), BLE Telemetry & Protocols, Sensor Data Acquisition, Edge AI Applications (Edge Impulse Studio).')
    add_bullet('Backend & Software Development: ', 'Node.js, NestJS, Express.js, JavaScript, Python, C / C++, PHP (CodeIgniter / Laravel), RESTful Web Services.')
    add_bullet('AI Tooling & Developer Productivity: ', 'ChatGPT, Gemini, Antigravity (AI-Assisted Code Generation, Rapid Prototyping & Workflow Acceleration).')
    add_bullet('Databases & DevOps Tools: ', 'MySQL, MongoDB, Relational Database Design, AWS S3, Postman, Git / GitLab, VS Code.')

    # Professional Experience
    add_sec_hdr('PROFESSIONAL EXPERIENCE', space_before=5)
    
    # Ingenious Technology Solutions (All 7 bullets cleanly on Page 1)
    add_p('Assistant Project Manager cum Software Engineer', bold=True, size=10.0, color='000000', space_before=2, space_after=1)
    add_p('Ingenious Technology Solutions Sdn. Bhd., Perai, Penang | Dec 2022 – Present', italic=True, size=9.0, color=SUBTEXT_COLOR, space_after=3)
    
    add_bullet('Enterprise Platform Leadership: ', 'Served as the primary technical Person-in-Charge (PIC) for a 14-module enterprise platform (MECHIS), proactively monitoring system health and coordinating with IT infrastructure teams to execute database capacity expansions, safeguarding continuous 24/7 uptime.')
    add_bullet('Embedded Touchscreen GUI & Hardware Integration: ', 'Designed an Embedded Touchscreen GUI for an ESP32-S3 AMOLED display using SquareLine Studio (LVGL framework), integrating touch interface components with underlying firmware logic and sensor data streaming pipelines.')
    add_bullet('Hardware-Software IoT Data Pipelines: ', 'Developed end-to-end telemetry workflows pairing edge microcontrollers (ESP32 / Arduino) with software interfaces via Bluetooth Low Energy (BLE), enabling real-time sensor data streaming and monitoring.')
    add_bullet('Embedded AI & Edge Application: ', 'Applied Edge AI models (Edge Impulse) for acoustic and signal classification, evaluating performance, signal processing, and low-power hardware integration.')
    add_bullet('Internal Tooling & System Automation: ', 'Architected and deployed a production-ready, role-based Ticketing Support System (Node.js/Express, MySQL) to automate incident management workflows and enforce SLA resolution timelines.')
    add_bullet('Python Data Diagnostics & AI Orchestration: ', 'Utilized Python scripts and AI-assisted developer workflows for serial telemetry logging, signal data analysis, and rapid system integration testing.')
    add_bullet('SDLC Governance & Requirement Analysis: ', 'Authored Functional Specification Documentation (FSD), evaluated Change Requests (CRs), and directed UAT cycles and client onboarding to maintain system stability.', space_after=0)

    # Page Break to start Page 2 cleanly
    doc.add_page_break()

    # ================= PAGE 2 =================
    add_sec_hdr('PROFESSIONAL EXPERIENCE (CONTINUED)', space_before=0)
    
    # Airei Sdn. Bhd.
    add_p('Software Engineer', bold=True, size=10.0, color='000000', space_before=3, space_after=1)
    add_p('Airei Sdn. Bhd., Cyberjaya, Selangor | Apr 2020 – Nov 2022', italic=True, size=9.0, color=SUBTEXT_COLOR, space_after=3)
    add_bullet('Application Development & Support: ', 'Executed full SDLC development for web applications using PHP (CodeIgniter) and JavaScript, adhering to strict client SLAs and system specifications.')
    add_bullet('Quality Assurance & Documentation: ', 'Conducted functional testing, authored technical end-user documentation, and supported production release deployments.')
    add_bullet('Client SLA & Incident Management: ', 'Managed Level 1 and Level 2 application support, troubleshooting software incidents and tracking resolution timelines to ensure strict adherence to service-level agreements.')

    # Majlis Bandaraya Ipoh
    add_p('Software Developer Intern', bold=True, size=10.0, color='000000', space_before=5, space_after=1)
    add_p('Majlis Bandaraya Ipoh, Perak | Feb 2019 – Jul 2019', italic=True, size=9.0, color=SUBTEXT_COLOR, space_after=3)
    add_bullet('Web Platform Development: ', 'Designed and built a web-based meeting scheduling platform (eJadual) using JavaScript, AJAX, and relational databases to automate municipal scheduling workflows.')

    # Education
    add_sec_hdr('EDUCATION', space_before=7)
    add_p('Bachelor’s Degree in Software Engineering', bold=True, size=10.0, color='000000', space_before=2, space_after=1)
    add_p('Universiti Putra Malaysia (UPM) | 2015 – 2020', italic=True, size=9.0, color=SUBTEXT_COLOR, space_after=3)
    add_bullet('Final Year Project: ', 'Two-Factor Authentication (2FA) Security Implementation for Internal Document Management Systems (Laravel / REST API).')

    # Certifications & Continuous Learning
    add_sec_hdr('CERTIFICATIONS & CONTINUOUS LEARNING', space_before=7)
    
    add_p('Edge AI Fundamentals | Edge Impulse (Aug 2026)', bold=True, size=9.5, color='000000', space_before=2, space_after=1)
    add_bullet('', 'Applied study in Edge AI architecture, signal processing, and deploying machine learning models on resource-constrained microcontrollers.', space_after=4, level=1)

    add_p('NestJS Enterprise Application & REST API Architecture | CodeSignal (Jan 2026)', bold=True, size=9.5, color='000000', space_before=3, space_after=1)
    add_bullet('', 'Structured backend application architecture, dependency injection, REST API development, and MongoDB integration.', space_after=4, level=1)

    add_p('Backend Development with Node.js & Express.js | Codefinity (Oct 2025)', bold=True, size=9.5, color='000000', space_before=3, space_after=1)
    add_bullet('', 'Server-side application logic, asynchronous routing, and database integrations.', space_after=4, level=1)

    doc.save(output_docx)
    print(f"Successfully generated: {output_docx}")

def create_resume_1page(output_docx):
    doc = docx.Document()
    
    # Page setup: A4, 0.42" top/bottom, 0.48" left/right
    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.top_margin = Inches(0.42)
        s.bottom_margin = Inches(0.42)
        s.left_margin = Inches(0.48)
        s.right_margin = Inches(0.48)

    base_pt = 8.5
    line_sp = 1.05

    def add_p(text='', bold=False, italic=False, size=base_pt, color='000000', align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=1.5):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_sp
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.italic = italic
            r.font.name = 'Arial'
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        return p

    def add_sec_hdr(title, space_before=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(title)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor.from_string(PRIMARY_COLOR)
        
        pPr = p._p.get_or_add_pPr()
        border_xml = f'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="4" w:space="1" w:color="{PRIMARY_COLOR}"/></w:pBdr>'
        pPr.append(parse_xml(border_xml))
        return p

    def add_bullet(prefix, text, space_after=1.2):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_sp
        p.paragraph_format.left_indent = Inches(0.14)
        
        rb = p.add_run('• ')
        rb.font.name = 'Arial'
        rb.font.size = Pt(base_pt)
        rb.font.color.rgb = RGBColor.from_string(PRIMARY_COLOR)
        
        if prefix:
            rp = p.add_run(prefix)
            rp.bold = True
            rp.font.name = 'Arial'
            rp.font.size = Pt(base_pt)
            rp.font.color.rgb = RGBColor.from_string('000000')
            
        rt = p.add_run(text)
        rt.bold = False
        rt.font.name = 'Arial'
        rt.font.size = Pt(base_pt)
        rt.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
        return p

    # Header
    add_p('NABIHAH BINTI TAQUDDIN AZMY', bold=True, size=13.5, color='000000', align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    add_p('Software Engineer | Embedded & Full-Stack Systems', bold=True, size=9.5, color=PRIMARY_COLOR, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    add_p('Perai, Penang | +60 13-519 8421 | nabihahtaqudin@gmail.com | linkedin.com/in/nabihah-taquddin-azmy', bold=False, size=8.0, color='404040', align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)

    # Professional Summary
    add_sec_hdr('PROFESSIONAL SUMMARY', space_before=2)
    add_p('Skilled Software Engineer and Technical Lead with over 5 years of professional experience building full-stack web applications, relational database infrastructure, and leading software platform delivery. Combines a strong software engineering foundation (Node.js, Express, NestJS, RESTful APIs, MySQL, MongoDB) with hands-on capabilities in embedded hardware integration (ESP32-S3 microcontrollers, Embedded Touchscreen GUIs, BLE telemetry, Edge AI). Experienced lead Person-in-Charge (PIC) for multi-module enterprise platforms, managing end-to-end SDLC workflows, Functional Specification Documentation (FSD), system monitoring, and cross-functional project execution.', size=base_pt, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2)

    # Core Competencies
    add_sec_hdr('CORE COMPETENCIES & TECH STACK', space_before=2)
    add_bullet('System Architecture & Integration: ', 'End-to-End System Design, Hardware-Software Integration, Embedded Touchscreen GUI Architecture, RESTful API Design, System Process Mapping, FSD Authoring & Technical Documentation.')
    add_bullet('IDEs, Embedded UI & Hardware Frameworks: ', 'ESP-IDF (C/C++), Arduino IDE, SquareLine Studio (LVGL Framework), Microcontrollers (ESP32-S3, Arduino), BLE Telemetry & Protocols, Sensor Data Acquisition, Edge AI Applications (Edge Impulse Studio).')
    add_bullet('Backend & Software Development: ', 'Node.js, NestJS, Express.js, JavaScript, Python, C / C++, PHP (CodeIgniter / Laravel), RESTful Web Services.')
    add_bullet('AI Tooling & Developer Productivity: ', 'ChatGPT, Gemini, Antigravity (AI-Assisted Code Generation, Rapid Prototyping & Workflow Acceleration).')
    add_bullet('Databases & DevOps Tools: ', 'MySQL, MongoDB, Relational Database Design, AWS S3, Postman, Git / GitLab, VS Code.')

    # Experience
    add_sec_hdr('PROFESSIONAL EXPERIENCE', space_before=2)
    
    # Ingenious
    add_p('Assistant Project Manager cum Software Engineer', bold=True, size=9.0, color='000000', space_before=1.5, space_after=0.5)
    add_p('Ingenious Technology Solutions Sdn. Bhd., Perai, Penang | Dec 2022 – Present', italic=True, size=8.0, color='333333', space_after=1.5)
    add_bullet('Enterprise Platform Leadership: ', 'Served as primary technical PIC for a 14-module enterprise platform (MECHIS), monitoring system health and coordinating with IT infrastructure teams to execute database expansions, safeguarding 24/7 uptime.')
    add_bullet('Embedded Touchscreen GUI & Hardware Integration: ', 'Designed Embedded Touchscreen GUI for ESP32-S3 AMOLED display via SquareLine Studio (LVGL), integrating touch UI with firmware logic and sensor data streaming pipelines.')
    add_bullet('Hardware-Software IoT Data Pipelines: ', 'Developed end-to-end telemetry workflows pairing edge MCUs (ESP32 / Arduino) with software interfaces via BLE, enabling real-time sensor data streaming and monitoring.')
    add_bullet('Embedded AI & Edge Application: ', 'Applied Edge AI models (Edge Impulse) for acoustic and signal classification, evaluating performance, signal processing, and low-power hardware integration.')
    add_bullet('Internal Tooling & System Automation: ', 'Architected and deployed a production-ready, role-based Ticketing Support System (Node.js/Express, MySQL) to automate incident management workflows and enforce SLA resolution timelines.')
    add_bullet('Python Data Diagnostics & AI Orchestration: ', 'Utilized Python scripts and AI-assisted developer workflows for serial telemetry logging, signal data analysis, and rapid system integration testing.')
    add_bullet('SDLC Governance & Requirement Analysis: ', 'Authored Functional Specification Documentation (FSD), evaluated Change Requests (CRs), and directed UAT cycles and client onboarding to maintain system stability.')

    # Airei
    add_p('Software Engineer', bold=True, size=9.0, color='000000', space_before=1.5, space_after=0.5)
    add_p('Airei Sdn. Bhd., Cyberjaya, Selangor | Apr 2020 – Nov 2022', italic=True, size=8.0, color='333333', space_after=1.5)
    add_bullet('Application Development & Support: ', 'Executed full SDLC development for web applications using PHP (CodeIgniter) and JavaScript, adhering to strict client SLAs and system specifications.')
    add_bullet('Quality Assurance & Documentation: ', 'Conducted functional testing, authored technical end-user documentation, and supported production release deployments.')

    # MBI
    add_p('Software Developer Intern', bold=True, size=9.0, color='000000', space_before=1.5, space_after=0.5)
    add_p('Majlis Bandaraya Ipoh, Perak | Feb 2019 – Jul 2019', italic=True, size=8.0, color='333333', space_after=1.5)
    add_bullet('Web Platform Development: ', 'Designed and built a web-based meeting scheduling platform (eJadual) using JavaScript, AJAX, and relational databases to automate municipal scheduling workflows.')

    # Education
    add_sec_hdr('EDUCATION', space_before=2)
    add_p('Bachelor’s Degree in Software Engineering | Universiti Putra Malaysia (UPM) | 2015 – 2020', bold=True, size=8.5, color='000000', space_before=1, space_after=1)
    add_bullet('Final Year Project: ', 'Two-Factor Authentication (2FA) Security Implementation for Internal Document Management Systems (Laravel / REST API).')

    # Certifications
    add_sec_hdr('CERTIFICATIONS & CONTINUOUS LEARNING', space_before=2)
    add_bullet('Edge AI Fundamentals | Edge Impulse (Aug 2026): ', 'Applied study in Edge AI architecture, signal processing, and deploying machine learning models on resource-constrained microcontrollers.')
    add_bullet('NestJS Enterprise Application & REST API Architecture | CodeSignal (Jan 2026): ', 'Structured backend application architecture, dependency injection, REST API development, and MongoDB integration.')
    add_bullet('Backend Development with Node.js & Express.js | Codefinity (Oct 2025): ', 'Server-side application logic, asynchronous routing, and database integrations.')

    doc.save(output_docx)
    print(f"Successfully generated: {output_docx}")

if __name__ == '__main__':
    create_resume_2page('Resume_Nabihah_Taquddin_Azmy_v6_2Page.docx')
    create_resume_1page('Resume_Nabihah_Taquddin_Azmy_v6_1Page.docx')
